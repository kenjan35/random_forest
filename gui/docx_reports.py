"""
Rapports de vol au format Word (.docx), modifiables, avec le même contenu que
les rapports PDF (gui/pdf_reports.py) : page de garde, récapitulatif, puis
détail de chaque vol avec la capture du graphique.
"""

import os
import tempfile
from datetime import datetime

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from gui.pdf_reports import _event_date, _event_time_range, _period_label, _safe_filename

ACCENT = "1CAAE2"
ZEBRA = "F1F5F9"
TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)
AUTHOR = "Fuel Theft Detector"


# ------------------------------------------------------------- mise en forme

def _shade(cell, hex_color):
    properties = cell._tc.get_or_add_tcPr()
    for old in properties.findall(qn("w:shd")):
        properties.remove(old)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    # l'ordre des éléments est imposé par le format Word : shd vient avant vAlign
    anchor = properties.find(qn("w:vAlign"))
    if anchor is not None:
        anchor.addprevious(shading)
    else:
        properties.append(shading)


def _repeat_header(row):
    """La ligne d'en-tête du tableau est répétée en haut de chaque page."""
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _cell_text(cell, text, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _table(document, headers, widths_mm):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for i, (cell, title) in enumerate(zip(table.rows[0].cells, headers)):
        _cell_text(cell, title, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(cell, ACCENT)
    _repeat_header(table.rows[0])
    _set_widths(table, widths_mm)
    return table


def _set_widths(table, widths_mm):
    """Largeurs fixes (Word et LibreOffice lisent la grille, pas seulement les cellules)."""
    table.autofit = False   # mise en page fixe (élément placé correctement par python-docx)
    for column, width in zip(table.columns, widths_mm):
        column.width = Mm(width)
    for grid_col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_mm):
        grid_col.set(qn("w:w"), str(int(Mm(width).twips)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_mm):
            cell.width = Mm(width)


def _add_row(table, values, widths_mm, zebra=False, bold=False, center_column=None):
    cells = table.add_row().cells
    for i, value in enumerate(values):
        align = WD_ALIGN_PARAGRAPH.CENTER if i == center_column else WD_ALIGN_PARAGRAPH.LEFT
        if value is not None:
            _cell_text(cells[i], value, bold=bold, align=align)
        if zebra:
            _shade(cells[i], ZEBRA)
    for cell, width in zip(cells, widths_mm):
        cell.width = Mm(width)
    return cells


def _page_number_field(paragraph):
    """Champ Word « numéro de page » (mis à jour automatiquement à l'ouverture)."""
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), tag)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)


def _setup_page(document):
    section = document.sections[0]
    section.page_height, section.page_width = Mm(297), Mm(210)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 22), ("Heading 1", 14)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    styles["Footer"].font.size = Pt(8)
    styles["Footer"].font.color.rgb = TEXT_MUTED
    zoom = document.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")   # attribut obligatoire absent du modèle par défaut

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.add_run(f"Rapport généré par {AUTHOR}\t\tPage ")
    _page_number_field(footer)


def _cover(document, site, period):
    for _ in range(4):
        document.add_paragraph()
    title = document.add_paragraph("RAPPORT VOL FUEL", style="Title")
    border = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "8")
        edge.set(qn("w:space"), "8")
        edge.set(qn("w:color"), "000000")
        border.append(edge)
    title._p.get_or_add_pPr().append(border)
    document.add_paragraph()
    for label, value in (("SITE : ", site), ("PERIODE : ", period)):
        paragraph = document.add_paragraph()
        paragraph.add_run(label).bold = True
        run = paragraph.add_run(value)
        run.bold = True
        for r in paragraph.runs:
            r.font.size = Pt(20)
    for _ in range(12):
        document.add_paragraph()
    document.add_paragraph(f"Date de création : {datetime.now():%d/%m/%Y}")
    document.add_paragraph(f"Auteur : {AUTHOR}")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ rapports

def create_site_report_docx(result, output_path, period_start, period_end, chart_saver):
    events = list(result.get("events", []))
    if not events:
        raise ValueError("Ce site ne contient aucun vol détecté.")

    site = str(result.get("site") or result.get("filename") or "Site")
    period = _period_label(period_start)
    document = Document()
    _setup_page(document)
    document.core_properties.title = f"RAPPORT VOL FUEL - {site}"
    document.core_properties.author = AUTHOR
    _cover(document, site, period)

    total = sum(float(event.get("fuel_drop") or 0) for event in events)

    document.add_paragraph("1. Récapitulatif", style="Heading 1")
    widths = (42, 62, 35)
    table = _table(document, ("Date", "Tranche horaire", "Quantité"), widths)
    for i, event in enumerate(events):
        _add_row(table, (_event_date(event), _event_time_range(event),
                         f"{float(event.get('fuel_drop') or 0):g} L"), widths, zebra=i % 2 == 1, center_column=2)
    _add_row(table, ("", "Total", f"{total:g} L"), widths, bold=True, center_column=2)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    document.add_paragraph(f"2. {period}", style="Heading 1")
    widths = (25, 27, 22, 96)
    table = _table(document, ("Date", "Tranche horaire", "Quantité", "Capture d'écran"), widths)
    temporary_images = []
    try:
        for i, event in enumerate(events):
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
                image_path = handle.name
            temporary_images.append(image_path)
            chart_saver(result, image_path, event)
            cells = _add_row(table, (_event_date(event), _event_time_range(event),
                                     f"{float(event.get('fuel_drop') or 0):g} L", None),
                             widths, zebra=i % 2 == 1, center_column=2)
            picture = cells[3].paragraphs[0]
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.add_run().add_picture(image_path, width=Mm(88))
        _add_row(table, ("", "Total", f"{total:g} L", ""), widths, bold=True, center_column=2)
        document.save(output_path)
    finally:
        for image_path in temporary_images:
            if os.path.exists(image_path):
                os.remove(image_path)


def create_docx_reports(results, output_dir, period_start, period_end, chart_saver):
    os.makedirs(output_dir, exist_ok=True)
    period = _period_label(period_start)
    paths = []
    for result in results:
        if result.get("error") or not result.get("has_theft"):
            continue
        site = _safe_filename(result.get("site") or result.get("filename"))
        path = os.path.join(output_dir, f"RAPPORT_VOL_FUEL_{site}_{_safe_filename(period)}.docx")
        create_site_report_docx(result, path, period_start, period_end, chart_saver)
        paths.append(path)
    return paths
