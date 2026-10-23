"""
Accès au stockage RMS SALTEC (MinIO, API compatible S3), sans dépendance Qt
pour pouvoir être testé seul. Suit le guide « RMS SALTEC Storage Platform –
Developer Integration Guide » v1.0 :
  - point de terminaison https://api-rms.appesca.com, bucket rms-saltec ;
  - droits en lecture seule : ListBucket et GetObject ;
  - arborescence pays / année / mois / jour, fichiers .json.gz ;
  - un seul client réutilisé, téléchargement en flux, réponses toujours fermées.
"""

import gzip
import json
import os
import re
from datetime import datetime

DEFAULT_ENDPOINT = "https://api-rms.appesca.com"
DEFAULT_BUCKET = "rms-saltec"
DEFAULT_REGION = "us-east-1"
KNOWN_COUNTRIES = ("CD", "MG", "TZ", "UG")   # dossiers racine cités dans le guide
ENV_ACCESS_KEY = "RMS_ACCESS_KEY"             # le guide recommande les variables d'environnement
ENV_SECRET_KEY = "RMS_SECRET_KEY"
PREVIEW_MAX_BYTES = 25 * 1024 * 1024
LIST_BATCH = 500


# ------------------------------------------------------------------ connexion

def normalize_endpoint(url):
    """'https://api-rms.appesca.com/' -> ('api-rms.appesca.com', True)."""
    text = str(url or "").strip()
    secure = not text.lower().startswith("http://")
    host = re.sub(r"^https?://", "", text, flags=re.IGNORECASE).strip().rstrip("/")
    if not host or "/" in host:
        raise ValueError("Point de terminaison invalide : indiquez seulement l'adresse du serveur, "
                         "par exemple https://api-rms.appesca.com")
    return host, secure


# Erreurs passagères de Cloudflare / du serveur : réessayées automatiquement.
RETRY_STATUS = (429, 500, 502, 503, 504, 520, 521, 522, 523, 524, 530)


def make_http_client(pool_size=64):
    """
    Client HTTP partagé : délais courts (l'interface ne reste pas bloquée
    plusieurs minutes), nouvelles tentatives sur les erreurs passagères, et
    prise en compte d'un proxy d'entreprise (variables HTTPS_PROXY / HTTP_PROXY).
    """
    import certifi
    import urllib3

    options = dict(
        timeout=urllib3.Timeout(connect=10, read=60),
        maxsize=pool_size,
        cert_reqs="CERT_REQUIRED",
        ca_certs=certifi.where(),
        retries=urllib3.Retry(total=3, backoff_factor=0.5, status_forcelist=RETRY_STATUS,
                              raise_on_status=False),
    )
    proxy = os.environ.get("HTTPS_PROXY") or os.environ.get("https_proxy") \
        or os.environ.get("HTTP_PROXY") or os.environ.get("http_proxy")
    return urllib3.ProxyManager(proxy, **options) if proxy else urllib3.PoolManager(**options)


def make_client(endpoint, access_key, secret_key, region=DEFAULT_REGION, pool_size=64):
    from minio import Minio

    host, secure = normalize_endpoint(endpoint)
    return Minio(host, access_key=access_key, secret_key=secret_key, secure=secure,
                 region=region or None, http_client=make_http_client(pool_size))


def check_endpoint(endpoint):
    """
    Diagnostic sans identifiants : interroge la sonde de santé MinIO.
    Retourne (ok, message). Permet de distinguer « serveur en panne »,
    « serveur injoignable » et « problème de clés ».
    """
    import urllib3

    try:
        host, secure = normalize_endpoint(endpoint)
    except ValueError as exc:
        return False, str(exc)
    url = f"{'https' if secure else 'http'}://{host}/minio/health/live"
    try:
        response = make_http_client(2).request("GET", url, retries=False, preload_content=True)
    except urllib3.exceptions.HTTPError as exc:
        return False, (f"{host} est injoignable depuis cet ordinateur ({type(exc).__name__}).\n"
                       "Vérifiez la connexion Internet, l'adresse, ou un éventuel proxy d'entreprise.")
    status = response.status
    if status == 200:
        return True, f"Le serveur {host} répond correctement (HTTP 200). En cas d'échec, vérifiez vos clés."
    body = response.data[:2000].decode("utf-8", errors="replace")
    described = describe_http_error(Exception(
        f"non-XML response from server; Response code: {status}, Body: {body}"), host)
    return False, described or f"Le serveur {host} a répondu HTTP {status}."


def env_credentials():
    return os.environ.get(ENV_ACCESS_KEY, ""), os.environ.get(ENV_SECRET_KEY, "")


# -------------------------------------------------------------------- listage

def entry_name(key, prefix):
    """Nom affiché d'un objet ou dossier, relatif au dossier courant."""
    return key[len(prefix):].rstrip("/") if key.startswith(prefix) else key.rstrip("/")


def _to_entry(item, prefix):
    key = str(item.object_name)
    is_dir = bool(getattr(item, "is_dir", False)) or key.endswith("/")
    modified = getattr(item, "last_modified", None)
    return {
        "key": key,
        "name": entry_name(key, prefix),
        "is_dir": is_dir,
        "size": None if is_dir else int(getattr(item, "size", 0) or 0),
        "modified": modified.astimezone().replace(tzinfo=None) if modified else None,
    }


def list_folder(client, bucket, prefix, should_stop=lambda: False, on_batch=None, batch_size=LIST_BATCH):
    """
    Liste UN niveau du dossier `prefix` (non récursif : les sous-dossiers
    apparaissent comme des entrées). Les résultats sont transmis par paquets à
    `on_batch` pour que l'interface se remplisse au fil du listage (un dossier
    « jour » peut contenir des dizaines de milliers de fichiers).
    Retourne le nombre total d'entrées.
    """
    batch, total = [], 0
    for item in client.list_objects(bucket, prefix=prefix, recursive=False):
        if should_stop():
            break
        entry = _to_entry(item, prefix)
        if not entry["name"]:
            continue
        batch.append(entry)
        total += 1
        if len(batch) >= batch_size:
            if on_batch:
                on_batch(batch)
            batch = []
    if batch and on_batch and not should_stop():
        on_batch(batch)
    return total


def is_access_denied(exc):
    return getattr(exc, "code", None) == "AccessDenied"


def list_root(client, bucket):
    """
    Dossiers à la racine du bucket. Le listage de la racine peut être refusé
    (AccessDenied) alors qu'il fonctionne avec un préfixe : on sonde alors les
    pays connus et on garde ceux qui répondent.
    Retourne (entrées, sondé) — `sondé` vaut True si la racine a été devinée.
    """
    entries = []
    try:
        list_folder(client, bucket, "", on_batch=entries.extend)
        return entries, False
    except Exception as exc:
        if not is_access_denied(exc):
            raise
    last_error = None
    for country in KNOWN_COUNTRIES:
        try:
            first = next(iter(client.list_objects(bucket, prefix=f"{country}/", recursive=False)), None)
        except Exception as exc:
            last_error = exc
            continue
        if first is not None:
            entries.append({"key": f"{country}/", "name": country, "is_dir": True,
                            "size": None, "modified": None})
    if not entries and last_error is not None:
        raise last_error
    return entries, True


# ------------------------------------------------------------------- contenu

def fetch_object(client, bucket, key, max_bytes=PREVIEW_MAX_BYTES):
    """Télécharge un objet en flux (réponse toujours fermée)."""
    response = client.get_object(bucket, key)
    try:
        chunks, size = [], 0
        for chunk in response.stream(1024 * 1024):
            size += len(chunk)
            if size > max_bytes:
                raise ValueError(f"Fichier trop volumineux pour l'aperçu (> {max_bytes // (1024 * 1024)} Mo).")
            chunks.append(chunk)
        return b"".join(chunks)
    finally:
        response.close()
        response.release_conn()


def decode_payload(raw):
    """
    Décompresse (si gzip) et met en forme le JSON.
    Retourne (texte_affiché, données_json_ou_None).
    """
    data = gzip.decompress(raw) if raw[:2] == b"\x1f\x8b" else raw
    text = data.decode("utf-8", errors="replace")
    try:
        parsed = json.loads(text)
    except ValueError:
        return text, None
    return json.dumps(parsed, indent=2, ensure_ascii=False), parsed


def payload_summary(parsed):
    """Résumé lisible d'un fichier SALTEC : mapping, siteId, nombre de mesures, période."""
    if not isinstance(parsed, dict):
        return ""
    parts = []
    for key in ("mapping", "siteId", "ingester"):
        if parsed.get(key):
            parts.append(str(parsed[key]))
    records = parsed.get("data")
    if isinstance(records, list):
        parts.append(f"{len(records)} mesure{'s' if len(records) > 1 else ''}")
        stamps = sorted(str(r.get("Timestamp")) for r in records if isinstance(r, dict) and r.get("Timestamp"))
        if stamps:
            parts.append(stamps[0].replace("T", " ") if len(stamps) == 1
                         else f"{stamps[0].replace('T', ' ')} → {stamps[-1].replace('T', ' ')}")
    return " · ".join(parts)


def format_size(size):
    if size is None:
        return ""
    for unit in ("o", "Ko", "Mo", "Go"):
        if size < 1024 or unit == "Go":
            return f"{size:.0f} {unit}" if unit == "o" else f"{size:.1f} {unit}"
        size /= 1024


def format_date(value):
    return value.strftime("%d/%m/%Y %H:%M") if isinstance(value, datetime) else ""


# -------------------------------------------------------------------- erreurs

# Codes de Cloudflare (le guide indique que l'API passe par un tunnel Cloudflare) :
# la requête n'a pas atteint le stockage, les clés ne sont pas en cause.
_CLOUDFLARE_STATUS = {
    520: "le serveur de stockage a renvoyé une réponse inattendue",
    521: "le serveur de stockage refuse la connexion (service arrêté ?)",
    522: "le serveur de stockage ne répond pas (délai de connexion dépassé)",
    523: "le serveur de stockage est injoignable depuis Cloudflare",
    524: "le serveur de stockage a mis trop de temps à répondre",
    525: "échec de la négociation SSL avec le serveur de stockage",
    526: "certificat SSL du serveur de stockage invalide",
    530: "Cloudflare n'arrive pas à joindre le serveur de stockage",
}
_CLOUDFLARE_CODES = {
    "1016": "nom DNS du serveur introuvable",
    "1033": "le tunnel Cloudflare vers le serveur est arrêté",
}
_S3_MESSAGES = {
    "InvalidAccessKeyId": "Clé d'accès invalide ou inconnue.",
    "SignatureDoesNotMatch": "Clé secrète incorrecte.",
    "AccessDenied": "Accès refusé : vos droits ne permettent pas de lister ou lire cet emplacement.",
    "NoSuchBucket": "Ce bucket n'existe pas ou ne vous est pas accessible.",
    "NoSuchKey": "Ce fichier n'existe plus.",
    "RequestTimeTooSkewed": "L'horloge de votre ordinateur est décalée : remettez-la à l'heure puis réessayez.",
}


def describe_http_error(exc, endpoint="api-rms.appesca.com"):
    """
    Message clair pour une réponse non-XML du serveur (page d'erreur HTML de
    Cloudflare ou d'un proxy au lieu d'une réponse S3). None sinon.
    """
    text = str(exc)
    status = getattr(exc, "_code", None) or getattr(exc, "status_code", None)
    if status is None:
        match = re.search(r"Response code:\s*(\d{3})", text)
        status = match.group(1) if match else None
    if status is None or ("non-XML" not in text and not str(status).startswith("5")):
        return None
    status = int(status)
    if status < 500:
        return None

    cause = _CLOUDFLARE_STATUS.get(status, "erreur côté serveur")
    cf_match = re.search(r"(?:error code:?|Error)\s*(1\d{3})", text, re.IGNORECASE)
    if cf_match:
        cf_code = cf_match.group(1)
        detail = f" : {_CLOUDFLARE_CODES[cf_code]}" if cf_code in _CLOUDFLARE_CODES else ""
        cause += f" (Cloudflare {cf_code}{detail})"
    return (
        f"Le serveur {endpoint} est indisponible (HTTP {status}) : {cause}.\n\n"
        "Vos clés d'accès ne sont pas en cause : la requête n'a pas atteint le stockage.\n"
        "Réessayez plus tard ou prévenez l'équipe BIG DATA & AI."
    )


def error_message(exc, endpoint="api-rms.appesca.com"):
    """Message en français pour n'importe quelle erreur d'accès au stockage."""
    code = getattr(exc, "code", None)
    if code in _S3_MESSAGES:
        return _S3_MESSAGES[code]
    http = describe_http_error(exc, endpoint)
    if http:
        return http
    name = type(exc).__name__
    if name in ("MaxRetryError", "NewConnectionError", "ConnectionError", "NameResolutionError") \
            or "Max retries exceeded" in str(exc):
        return (f"Impossible de joindre {endpoint}. Vérifiez votre connexion Internet "
                "et l'adresse du point de terminaison.")
    if isinstance(exc, ValueError):
        return str(exc)
    if code:
        return f"Erreur S3 ({code}) : {getattr(exc, 'message', '') or exc}"
    return f"{name} : {exc}"


# --------------------------------------------------------- extraction en lot

def day_prefixes(country, start, end):
    """['MG/2026/01/01/', ...] pour chaque jour de [start, end] (dates incluses)."""
    from datetime import timedelta

    if end < start:
        start, end = end, start
    days = (end - start).days + 1
    return [f"{country}/{(start + timedelta(days=i)):%Y/%m/%d}/" for i in range(days)]


def output_path(output_dir, key, decompress):
    """Chemin local qui reproduit l'arborescence du bucket (…/MG/2026/01/01/fichier.json)."""
    relative = key[:-3] if decompress and key.endswith(".gz") else key
    parts = [p for p in relative.split("/") if p]
    if any(p in ("..", ".") for p in parts):
        raise ValueError(f"Nom d'objet invalide : {key}")
    return os.path.join(output_dir, *parts)


INDEX_COLUMNS = ["key", "mapping", "siteId", "ingester", "mesures", "debut", "fin", "statut", "fichier_local"]


def _index_row(key, parsed, status, path=""):
    parsed = parsed if isinstance(parsed, dict) else {}
    records = parsed.get("data") if isinstance(parsed.get("data"), list) else []
    stamps = sorted(str(r.get("Timestamp")) for r in records if isinstance(r, dict) and r.get("Timestamp"))
    return {
        "key": key, "mapping": parsed.get("mapping", ""), "siteId": parsed.get("siteId", ""),
        "ingester": parsed.get("ingester", ""), "mesures": len(records),
        "debut": stamps[0] if stamps else "", "fin": stamps[-1] if stamps else "",
        "statut": status, "fichier_local": path,
    }


def _process_object(client, bucket, key, output_dir, decompress, mappings, skip_existing):
    """Télécharge un objet et l'enregistre s'il correspond aux mappings demandés."""
    path = output_path(output_dir, key, decompress)
    if skip_existing and os.path.exists(path):
        return _index_row(key, None, "deja_present", path), 0
    raw = fetch_object(client, bucket, key, max_bytes=2 ** 40)
    text, parsed = decode_payload(raw)
    if mappings:
        mapping = str((parsed or {}).get("mapping", "")).strip().upper() if isinstance(parsed, dict) else ""
        if mapping not in mappings:
            return _index_row(key, parsed, "ignore_mapping"), len(raw)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    data = (gzip.decompress(raw) if raw[:2] == b"\x1f\x8b" else raw) if decompress else raw
    tmp = path + ".part"
    with open(tmp, "wb") as handle:   # écriture atomique : pas de fichier à moitié écrit si on arrête
        handle.write(data)
    os.replace(tmp, path)
    return _index_row(key, parsed, "telecharge", path), len(raw)


def extract_range(client, bucket, prefixes, output_dir, decompress=True, mappings=None,
                  workers=8, skip_existing=True, should_stop=lambda: False,
                  on_progress=None, index_path=None):
    """
    Extrait tous les fichiers des dossiers `prefixes` (un par jour) vers
    `output_dir`, en parallèle. `mappings` (ensemble de codes, ex. {"TATS032"})
    ne garde que les fichiers de ces sites. Un index CSV (un fichier = une
    ligne : mapping, siteId, période, statut) est écrit au fil de l'eau.
    Les callbacks sont appelés depuis CE thread uniquement.
    Retourne les statistiques finales.
    """
    import csv
    from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait

    mappings = {m.strip().upper() for m in mappings or [] if m.strip()} or None
    stats = {"listes": 0, "telecharges": 0, "deja_presents": 0, "ignores": 0, "erreurs": 0,
             "octets": 0, "jour": "", "jours_faits": 0, "jours_total": len(prefixes), "derniere_erreur": ""}
    os.makedirs(output_dir, exist_ok=True)
    index_file = open(index_path, "w", newline="", encoding="utf-8-sig") if index_path else None
    writer = csv.DictWriter(index_file, fieldnames=INDEX_COLUMNS) if index_file else None
    if writer:
        writer.writeheader()

    def report():
        if on_progress:
            on_progress(dict(stats))

    def collect(done_futures, pending):
        for future in done_futures:
            key = pending.pop(future)
            try:
                row, size = future.result()
            except Exception as exc:
                stats["erreurs"] += 1
                stats["derniere_erreur"] = f"{key} : {error_message(exc)}"
                row, size = _index_row(key, None, "erreur"), 0
            stats["octets"] += size
            counter = {"telecharge": "telecharges", "deja_present": "deja_presents",
                       "ignore_mapping": "ignores"}.get(row["statut"])
            if counter:
                stats[counter] += 1
            if writer:
                writer.writerow(row)

    pending = {}
    try:
        with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
            for prefix in prefixes:
                if should_stop():
                    break
                stats["jour"] = prefix
                report()
                for item in client.list_objects(bucket, prefix=prefix, recursive=True):
                    if should_stop():
                        break
                    key = str(item.object_name)
                    if key.endswith("/"):
                        continue
                    stats["listes"] += 1
                    pending[pool.submit(_process_object, client, bucket, key, output_dir,
                                        decompress, mappings, skip_existing)] = key
                    # limite la file d'attente : la mémoire reste stable même sur 100 000 fichiers
                    while len(pending) >= workers * 4:
                        done, _ = wait(pending, return_when=FIRST_COMPLETED)
                        collect(done, pending)
                        report()
                stats["jours_faits"] += 1
            if should_stop():
                for future in pending:
                    future.cancel()
            done, _ = wait(pending)
            collect([f for f in done if not f.cancelled()], pending)
    finally:
        if index_file:
            index_file.close()
    report()
    return stats


def read_mappings_from_excel(path):
    """Codes de la colonne « mapping » d'un fichier Excel (cherchée dans les 30 premières lignes)."""
    import unicodedata
    from openpyxl import load_workbook

    def norm(value):
        return unicodedata.normalize("NFKC", str(value)).replace("\ufeff", "").replace("\xa0", " ").strip().casefold()

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        cell = next((c for row in sheet.iter_rows(min_row=1, max_row=30) for c in row
                     if c.value is not None and norm(c.value) == "mapping"), None)
        if cell is None:
            raise ValueError("La colonne 'mapping' est introuvable dans les 30 premières lignes.")
        codes = []
        for (value,) in sheet.iter_rows(min_row=cell.row + 1, min_col=cell.column,
                                        max_col=cell.column, values_only=True):
            code = str(value).strip().upper() if value is not None else ""
            if code and code not in codes:
                codes.append(code)
        return codes
    finally:
        workbook.close()


def parse_mappings(text):
    """'TATS032, tanj044 ; MGA012' -> ['TATS032', 'TANJ044', 'MGA012']."""
    codes = []
    for code in re.split(r"[\s,;]+", str(text or "")):
        code = code.strip().upper()
        if code and code not in codes:
            codes.append(code)
    return codes


# ------------------------------------------------- téléchargement parallèle

DOWNLOAD_MAX_WORKERS = 64
MAX_CONSECUTIVE_ERRORS = 50   # au-delà, le serveur est probablement indisponible : on arrête


def expand_prefixes(client, bucket, prefixes, should_stop=lambda: False):
    """
    Découpe chaque dossier en sous-dossiers (un niveau) pour pouvoir les lister en
    parallèle : « MG/2026/01/ » devient « MG/2026/01/01/ », « …/02/ », etc.
    Retourne (préfixes à lister récursivement, clés de fichiers trouvées au passage).
    """
    shards, keys = [], []
    for prefix in prefixes:
        children = []
        for item in client.list_objects(bucket, prefix=prefix, recursive=False):
            if should_stop():
                break
            name = str(item.object_name)
            if name.endswith("/") or getattr(item, "is_dir", False):
                children.append(name if name.endswith("/") else name + "/")
            else:
                keys.append(name)
        shards.extend(children)
    return shards, keys


def fast_download(client, bucket, output_dir, base_prefix="", keys=None, prefixes=None,
                  workers=32, list_workers=8, zip_path=None, decompress=False, skip_existing=True,
                  should_stop=lambda: False, on_progress=None, progress_every=0.5):
    """
    Télécharge beaucoup de petits fichiers vite :
      - les dossiers sont listés en parallèle (un fil par sous-dossier/jour) ;
      - les fichiers sont téléchargés par `workers` fils simultanés, dès qu'ils
        sont listés (pas besoin d'attendre la fin du listage) ;
      - une seule connexion HTTPS réutilisée par fil (pas de nouvelle poignée
        de main TLS à chaque fichier) ;
      - option `zip_path` : tout est écrit dans UNE archive ZIP au lieu de
        250 000 fichiers (beaucoup plus rapide sous Windows : l'antivirus et le
        système de fichiers ralentissent fortement la création de petits fichiers).
    `keys` (sélection) et/ou `prefixes` (dossiers) indiquent quoi télécharger ;
    les chemins sont rendus relatifs à `base_prefix`.
    Les callbacks sont appelés depuis le fil appelant uniquement.
    """
    import threading
    import time
    import zipfile
    from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
    from queue import Empty, Queue

    workers = max(1, min(int(workers), DOWNLOAD_MAX_WORKERS))
    stats = {"listes": 0, "telecharges": 0, "deja_presents": 0, "erreurs": 0, "octets": 0,
             "listage_termine": False, "derniere_erreur": "", "debut": time.time()}
    key_queue = Queue(maxsize=workers * 8)
    listing_error = []
    DONE = object()
    # arrêt demandé par l'utilisateur OU fin anticipée (erreur) : débloque aussi le listage
    abort = threading.Event()
    user_stop = should_stop
    should_stop = lambda: abort.is_set() or user_stop()  # noqa: E731

    def relative(key):
        rel = key[len(base_prefix):] if base_prefix and key.startswith(base_prefix) else key
        if decompress and rel.endswith(".gz"):
            rel = rel[:-3]
        parts = [p for p in rel.split("/") if p]
        if not parts or any(p in ("..", ".") for p in parts):
            raise ValueError(f"Nom d'objet invalide : {key}")
        return "/".join(parts)

    def put(item):
        """Ajoute à la file sans rester bloqué si l'utilisateur arrête."""
        from queue import Full
        while True:
            try:
                key_queue.put(item, timeout=0.2)
                return True
            except Full:
                if should_stop():
                    return False

    # ---- listage (fil producteur) : sous-dossiers listés en parallèle
    def producer():
        try:
            for key in keys or []:
                if should_stop() or not put(key):
                    return
            if prefixes:
                shards, direct = expand_prefixes(client, bucket, prefixes, should_stop)
                for key in direct:
                    if not put(key):
                        return

                def list_shard(prefix):
                    for item in client.list_objects(bucket, prefix=prefix, recursive=True):
                        if should_stop():
                            return
                        name = str(item.object_name)
                        if not name.endswith("/") and not put(name):
                            return

                with ThreadPoolExecutor(max_workers=max(1, list_workers)) as pool:
                    for future in [pool.submit(list_shard, p) for p in shards]:
                        future.result()
        except Exception as exc:
            listing_error.append(exc)
        finally:
            put(DONE)

    def download(key):
        rel = relative(key)
        target = os.path.join(output_dir, *rel.split("/")) if zip_path is None else None
        if target is not None and skip_existing and os.path.exists(target):
            return key, rel, None, "deja_present"
        raw = fetch_object(client, bucket, key, max_bytes=2 ** 40)
        data = (gzip.decompress(raw) if raw[:2] == b"\x1f\x8b" else raw) if decompress else raw
        if target is not None:
            os.makedirs(os.path.dirname(target), exist_ok=True)
            tmp = target + ".part"
            with open(tmp, "wb") as handle:
                handle.write(data)
            os.replace(tmp, target)
            return key, rel, len(raw), "telecharge"
        return key, rel, data, "telecharge"   # écrit dans le ZIP par le fil appelant

    archive = None
    if zip_path is not None:
        os.makedirs(os.path.dirname(os.path.abspath(zip_path)) or ".", exist_ok=True)
        # les .gz sont déjà compressés : on les range tels quels (plus rapide)
        archive = zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED if decompress else zipfile.ZIP_STORED,
                                  allowZip64=True)
    else:
        os.makedirs(output_dir, exist_ok=True)

    threading.Thread(target=producer, daemon=True).start()
    pending, consecutive_errors, last_report = set(), 0, 0.0

    def report(force=False):
        nonlocal last_report
        now = time.time()
        if on_progress and (force or now - last_report >= progress_every):
            last_report = now
            stats["duree"] = now - stats["debut"]
            on_progress(dict(stats))

    def collect(done):
        nonlocal consecutive_errors
        for future in done:
            pending.discard(future)
            try:
                key, rel, payload, status = future.result()
            except Exception as exc:
                stats["erreurs"] += 1
                stats["derniere_erreur"] = error_message(exc)
                consecutive_errors += 1
                if consecutive_errors >= MAX_CONSECUTIVE_ERRORS:
                    raise RuntimeError(f"{MAX_CONSECUTIVE_ERRORS} échecs d'affilée, téléchargement arrêté.\n"
                                       f"Dernière erreur : {stats['derniere_erreur']}")
                continue
            consecutive_errors = 0
            if status == "deja_present":
                stats["deja_presents"] += 1
                continue
            if archive is not None:
                archive.writestr(rel, payload)
                stats["octets"] += len(payload)
            else:
                stats["octets"] += payload
            stats["telecharges"] += 1

    try:
        with ThreadPoolExecutor(max_workers=workers) as pool:
            finished_listing = False
            while not finished_listing or pending:
                if should_stop():
                    break
                # remplit la file de téléchargements sans dépasser workers × 4 en attente
                while not finished_listing and len(pending) < workers * 4:
                    try:
                        key = key_queue.get(timeout=0.05)
                    except Empty:
                        break
                    if key is DONE:
                        finished_listing = True
                        stats["listage_termine"] = True
                        break
                    stats["listes"] += 1
                    pending.add(pool.submit(download, key))
                if pending:
                    done, _ = wait(pending, timeout=0.2, return_when=FIRST_COMPLETED)
                    collect(done)
                report()
            if should_stop():
                for future in pending:
                    future.cancel()
                done, _ = wait(pending)
                collect([f for f in done if not f.cancelled()])
    finally:
        abort.set()   # libère le fil de listage s'il attend encore de la place dans la file
        if archive is not None:
            archive.close()
    if listing_error and not user_stop():
        stats["derniere_erreur"] = f"Listage interrompu : {error_message(listing_error[0])}"
        stats["erreurs"] += 1
    stats["arrete"] = bool(user_stop())
    report(force=True)
    return stats