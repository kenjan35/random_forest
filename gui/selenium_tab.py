import os
import unicodedata
from datetime import datetime

from openpyxl import load_workbook
from PySide6.QtCore import QDate, Qt, Signal
from PySide6.QtGui import QIcon, QCursor
from PySide6.QtWidgets import (
    QDateEdit, QFormLayout, QGroupBox, QHBoxLayout, QHeaderView, QLabel,
    QAbstractItemView, QLineEdit, QMessageBox, QPushButton, QTableWidget,
    QTableWidgetItem, QVBoxLayout, QWidget, QFileDialog, QTabWidget, QSplitter,
)

from gui.chart import FuelChart, THEMES, annotate_drop
from gui.export import write_events_csv
from gui.cursor_utils import apply_pointer_cursor, set_button_enabled, set_button_busy
from gui.pdf_reports import create_reports
from gui.docx_reports import create_docx_reports
from gui.site_history import record_extraction
from gui.workers import SeleniumExtractionWorker


class SeleniumTab(QWidget):
    history_updated = Signal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self.worker = None
        self.results = []
        self.sites = []
        self._sites_requested = 0
        self._build_ui()

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(20, 20, 20, 20)
        root.setSpacing(14)

        title = QLabel("Extraction Power-Noc")
        title.setObjectName("pageTitle")
        root.addWidget(title)
        subtitle = QLabel("Connectez-vous à TOA, choisissez une période, puis récupérez les mesures directement.")
        subtitle.setObjectName("mutedLabel")
        root.addWidget(subtitle)

        self.pages = QTabWidget()
        extraction_page = QWidget()
        extraction_layout = QVBoxLayout(extraction_page)
        extraction_layout.setContentsMargins(10, 12, 10, 10)
        extraction_layout.setSpacing(14)
        results_page = QWidget()
        results_layout = QVBoxLayout(results_page)
        results_layout.setContentsMargins(10, 12, 10, 10)
        results_layout.setSpacing(10)
        self.pages.addTab(extraction_page, "Extraction")
        self.pages.addTab(results_page, "Résultats")
        self.pages.tabBar().setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        root.addWidget(self.pages, stretch=1)

        form_box = QGroupBox("Connexion")
        form = QFormLayout(form_box)
        form.setContentsMargins(16, 18, 16, 16)
        form.setHorizontalSpacing(18)
        form.setVerticalSpacing(12)
        self.url_edit = QLineEdit("https://toa.power-noc.net/")
        self.user_edit = QLineEdit()
        self.user_edit.setPlaceholderText("Identifiant TOA")
        self.password_edit = QLineEdit()
        self.password_edit.setPlaceholderText("Mot de passe")
        self.password_edit.setEchoMode(QLineEdit.EchoMode.Password)
        self.password_visibility_action = self.password_edit.addAction(
            QIcon("assets/icons/eye.svg"), QLineEdit.ActionPosition.TrailingPosition
        )
        self.password_visibility_action.setToolTip("Afficher le mot de passe")
        self.password_visibility_action.triggered.connect(self.toggle_password_visibility)
        self.start_date = QDateEdit(QDate.currentDate().addMonths(-1))
        self.end_date = QDateEdit(QDate.currentDate())
        for date_edit in (self.start_date, self.end_date):
            date_edit.setCalendarPopup(True)
            date_edit.setDisplayFormat("dd/MM/yyyy")
        form.addRow("URL", self.url_edit)
        form.addRow("Identifiant", self.user_edit)
        form.addRow("Mot de passe", self.password_edit)
        extraction_layout.addWidget(form_box)

        site_box = QGroupBox("Sites à extraire")
        site_layout = QVBoxLayout(site_box)
        site_actions = QHBoxLayout()
        self.excel_btn = QPushButton("Charger un fichier Excel")
        self.excel_btn.clicked.connect(self.load_sites_file)
        apply_pointer_cursor(self.excel_btn)
        self.select_sites_btn = QPushButton("Tout sélectionner")
        self.select_sites_btn.setObjectName("secondary")
        self.select_sites_btn.setEnabled(False)
        self.select_sites_btn.clicked.connect(self.select_all_sites)
        apply_pointer_cursor(self.select_sites_btn)
        self.site_file_label = QLabel("Aucun fichier Excel sélectionné")
        self.site_file_label.setObjectName("mutedLabel")
        site_actions.addWidget(self.excel_btn)
        site_actions.addWidget(self.select_sites_btn)
        site_actions.addWidget(QLabel("Période"))
        site_actions.addWidget(self.start_date)
        site_actions.addWidget(QLabel(" au "))
        site_actions.addWidget(self.end_date)
        site_actions.addWidget(self.site_file_label, stretch=1)
        site_layout.addLayout(site_actions)
        self.site_table = QTableWidget(0, 2)
        self.site_table.setHorizontalHeaderLabels(["Site", "Sélection"])
        self.site_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.site_table.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.site_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.site_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.site_table.horizontalHeader().setStretchLastSection(True)
        self.site_table.itemSelectionChanged.connect(self.update_site_count)
        site_layout.addWidget(self.site_table)
        extraction_layout.addWidget(site_box, stretch=1)

        actions = QHBoxLayout()
        self.extract_btn = QPushButton("Lancer l'extraction")
        self.extract_btn.clicked.connect(self.start_extraction)
        apply_pointer_cursor(self.extract_btn)
        self.status_label = QLabel("Prêt")
        self.status_label.setObjectName("mutedLabel")
        actions.addWidget(self.extract_btn)
        actions.addWidget(self.status_label, stretch=1)
        extraction_layout.addLayout(actions)

        self.table = QTableWidget(0, 7)
        self.table.setHorizontalHeaderLabels([
            "Site Excel", "Statut", "Nb périodes", "Début alerte", "Fin alerte", "Total loss", "Tank capacity"
        ])
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setColumnWidth(3, 145)
        self.table.setColumnWidth(4, 145)
        self.table.itemSelectionChanged.connect(self.on_row_selected)
        self.chart = FuelChart()
        results_splitter = QSplitter(Qt.Orientation.Vertical)
        results_splitter.addWidget(self.table)
        results_splitter.addWidget(self.chart)
        results_splitter.setStretchFactor(0, 1)
        results_splitter.setStretchFactor(1, 2)
        results_actions = QHBoxLayout()
        self.export_btn = QPushButton("Exporter les données")
        self.export_btn.setObjectName("secondary")
        self.export_btn.setEnabled(False)
        self.export_btn.clicked.connect(self.export_csv)
        apply_pointer_cursor(self.export_btn)
        self.pdf_btn = QPushButton("Générer les rapports PDF")
        self.pdf_btn.setObjectName("secondary")
        self.pdf_btn.setEnabled(False)
        self.pdf_btn.clicked.connect(self.generate_pdf_reports)
        apply_pointer_cursor(self.pdf_btn)
        self.docx_btn = QPushButton("Générer les rapports Word")
        self.docx_btn.setObjectName("secondary")
        self.docx_btn.setToolTip("Rapports .docx modifiables (même contenu que les PDF)")
        self.docx_btn.setEnabled(False)
        self.docx_btn.clicked.connect(self.generate_docx_reports)
        apply_pointer_cursor(self.docx_btn)
        results_actions.addWidget(self.export_btn)
        results_actions.addWidget(self.pdf_btn)
        results_actions.addWidget(self.docx_btn)
        results_actions.addStretch()
        results_layout.addLayout(results_actions)
        results_layout.addWidget(results_splitter, stretch=1)

    def toggle_password_visibility(self):
        is_visible = self.password_edit.echoMode() == QLineEdit.EchoMode.Password
        self.password_edit.setEchoMode(
            QLineEdit.EchoMode.Normal if is_visible else QLineEdit.EchoMode.Password
        )
        self.password_visibility_action.setIcon(
            QIcon("assets/icons/eye-off.svg" if is_visible else "assets/icons/eye.svg")
        )
        self.password_visibility_action.setToolTip(
            "Masquer le mot de passe" if is_visible else "Afficher le mot de passe"
        )

    def start_extraction(self):
        if not self.user_edit.text().strip() or not self.password_edit.text():
            QMessageBox.warning(self, "Informations manquantes", "Saisissez votre identifiant et votre mot de passe.")
            return
        if self.start_date.date() > self.end_date.date():
            QMessageBox.warning(self, "Période invalide", "La date de début doit précéder la date de fin.")
            return
        selected_rows = sorted({index.row() for index in self.site_table.selectionModel().selectedRows()})
        sites = [self.site_table.item(row, 0).text() for row in selected_rows]
        if not sites:
            QMessageBox.warning(self, "Sites manquants", "Chargez un fichier Excel et sélectionnez au moins un site.")
            return
        self.table.setRowCount(0)
        self.results = []
        self._sites_requested = len(sites)
        set_button_busy(self.extract_btn, True)
        set_button_enabled(self.extract_btn, False)
        set_button_enabled(self.export_btn, False)
        set_button_enabled(self.pdf_btn, False)
        set_button_enabled(self.docx_btn, False)
        self.status_label.setText(f"0 / {len(sites)} sites traités")
        self.worker = SeleniumExtractionWorker(
            self.url_edit.text().strip(), self.user_edit.text(), self.password_edit.text(),
            self.start_date.date().toString("yyyy-MM-dd"), self.end_date.date().toString("yyyy-MM-dd"),
            sites,
        )
        self.worker.result_ready.connect(self.on_result_ready)
        self.worker.finished_ok.connect(self.on_success)
        self.worker.finished_error.connect(self.on_error)
        self.worker.dates_adjusted.connect(self.on_dates_adjusted)
        self.worker.start()

    def on_dates_adjusted(self, start_date, end_date):
        self._update_site_progress()

    def on_result_ready(self, result):
        self.results.append(result)
        self._append_result_to_table(result)
        try:
            record_extraction(
                [result],
                self.start_date.date().toString("yyyy-MM-dd"),
                self.end_date.date().toString("yyyy-MM-dd"),
            )
            self.history_updated.emit()
        except Exception as exc:  # l'historique ne doit jamais bloquer l'extraction
            self.status_label.setText(f"Historique non enregistré : {exc}")
        set_button_enabled(self.export_btn, True)
        set_button_enabled(self.pdf_btn, any(item.get("has_theft") for item in self.results))
        set_button_enabled(self.docx_btn, any(item.get("has_theft") for item in self.results))
        self._update_site_progress()
        self.pages.setCurrentIndex(1)

    def _update_site_progress(self):
        # compte les sites distincts (un site peut produire plusieurs résultats)
        done = len({result.get("site") for result in self.results})
        self.status_label.setText(f"{done} / {self._sites_requested} sites traités")

    def _append_result_to_table(self, result):
        row_index = self.table.rowCount()
        self.table.insertRow(row_index)
        site = result.get("site", "")
        file_item = QTableWidgetItem(site)
        file_item.setData(Qt.ItemDataRole.UserRole, result)
        self.table.setItem(row_index, 0, file_item)
        if result.get("error"):
            status, events, start_alert, end_alert, total_loss, tank_capacity = "Erreur", "-", "", "", "", ""
        elif result.get("has_theft"):
            status = "Suspect"
            events = result.get("n_events", 0)
            alert_events = result.get("events", [])
            start_alert, end_alert = self._alert_range(alert_events)
            total_loss = self._total_loss(alert_events)
            tank_capacity = self._format_capacity(result.get("tank_capacity"))
        elif result.get("has_sensor_glitch"):
            status = "Data incomplete"
            events = result.get("n_events_ignored", 0)
            alert_events = result.get("events_ignored", [])
            start_alert, end_alert = self._alert_range(alert_events)
            total_loss = ""
            tank_capacity = self._format_capacity(result.get("tank_capacity"))
        else:
            status, events, start_alert, end_alert, total_loss = "Normal", 0, "", "", "0 L"
            tank_capacity = self._format_capacity(result.get("tank_capacity"))
        self.table.setItem(row_index, 1, QTableWidgetItem(status))
        self.table.setItem(row_index, 2, QTableWidgetItem(str(events)))
        self.table.setItem(row_index, 3, QTableWidgetItem(start_alert))
        self.table.setItem(row_index, 4, QTableWidgetItem(end_alert))
        self.table.setItem(row_index, 5, QTableWidgetItem(total_loss))
        self.table.setItem(row_index, 6, QTableWidgetItem(tank_capacity))

    @staticmethod
    def _format_alert_timestamp(value):
        text = str(value or "").strip()
        for format_string in (
            "%Y-%m-%dT%H:%M:%S.%f",
            "%Y-%m-%dT%H:%M:%S",
            "%Y-%m-%d %H:%M:%S.%f",
            "%Y-%m-%d %H:%M:%S",
        ):
            try:
                return datetime.strptime(text.rstrip("Z"), format_string).strftime("%Y-%m-%d %H:%M")
            except ValueError:
                continue
        return text[:16]

    @classmethod
    def _alert_range(cls, events):
        if not events:
            return "", ""
        return (
            cls._format_alert_timestamp(events[0].get("ts_start")),
            cls._format_alert_timestamp(events[-1].get("ts_end")),
        )

    @staticmethod
    def _total_loss(events):
        total = sum(float(event.get("fuel_drop") or 0) for event in events)
        return f"{total:.1f} L"

    @staticmethod
    def _format_capacity(value):
        if value is None or value == "":
            return ""
        return f"{float(value):.1f} L"

    def load_sites_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Choisir le fichier Excel des sites", "", "Excel (*.xlsx *.xls)")
        if not path:
            return
        try:
            # Lecture brute pour conserver exactement le texte du fichier et
            # éviter qu'une ligne de métadonnées soit prise pour l'en-tête.
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sites = self._read_sites(workbook.active)
            finally:
                workbook.close()  # sinon le fichier reste verrouillé sous Windows
        except Exception as exc:
            QMessageBox.critical(self, "Fichier Excel", f"Impossible de lire les sites : {exc}")
            return
        self._show_sites(path, sites)

    def _read_sites(self, sheet):
        """Lit la colonne 'mapping' (cherchée dans les 30 premières lignes)."""
        mapping_cell = next(
            (cell for row in sheet.iter_rows(min_row=1, max_row=30)
             for cell in row if self._normalize_excel_value(cell.value) == "mapping"),
            None,
        )
        if mapping_cell is None:
            raise ValueError("La colonne 'mapping' est introuvable dans les 30 premières lignes.")

        site_column = mapping_cell.column
        sites = []
        seen = set()
        for (value,) in sheet.iter_rows(
            min_row=mapping_cell.row + 1,
            min_col=site_column,
            max_col=site_column,
            values_only=True,
        ):
            if value is None:
                continue
            site = str(value).strip()
            key = self._normalize_excel_value(site)
            if site and key not in seen:
                sites.append(site)
                seen.add(key)
        if not sites:
            raise ValueError("Aucun nom de site exploitable n'a été trouvé.")
        return sites

    def _show_sites(self, path, sites):
        self.sites = sites
        self.site_table.setRowCount(len(sites))
        for row, site in enumerate(sites):
            self.site_table.setItem(row, 0, QTableWidgetItem(site))
            self.site_table.setItem(row, 1, QTableWidgetItem("Disponible"))
        self.site_file_label.setText(f"{os.path.basename(path)} : {len(sites)} site(s) unique(s)")
        set_button_enabled(self.select_sites_btn, True)

    @staticmethod
    def _normalize_excel_value(value):
        text = unicodedata.normalize("NFKC", str(value))
        return text.replace("\ufeff", "").replace("\xa0", " ").strip().casefold()

    def select_all_sites(self):
        self.site_table.selectAll()
        self.update_site_count()

    def update_site_count(self):
        count = len(self.site_table.selectionModel().selectedRows())
        if self.sites:
            self.site_file_label.setText(f"{count}/{len(self.sites)} site(s) sélectionné(s)")

    def on_success(self, results):
        # chaque résultat a déjà été enregistré dans on_result_ready
        set_button_busy(self.extract_btn, False)
        set_button_enabled(self.extract_btn, True)
        set_button_enabled(self.export_btn, bool(results))
        set_button_enabled(self.pdf_btn, any(result.get("has_theft") for result in results))
        set_button_enabled(self.docx_btn, any(result.get("has_theft") for result in results))
        self._update_site_progress()
        self.pages.setCurrentIndex(1)
        if results:
            self.table.selectRow(0)

    def on_row_selected(self):
        selected = self.table.selectionModel().selectedRows()
        if not selected:
            self.chart.clear()
            return
        item = self.table.item(selected[0].row(), 0)
        result = item.data(Qt.ItemDataRole.UserRole) if item else None
        if result and not result.get("error"):
            if result.get("fuel") is not None:
                self.chart.plot(result)
        else:
            self.chart.clear()

    def on_error(self, message):
        set_button_busy(self.extract_btn, False)
        set_button_enabled(self.extract_btn, True)
        if self.results:
            set_button_enabled(self.export_btn, True)
            set_button_enabled(self.pdf_btn, any(result.get("has_theft") for result in self.results))
            set_button_enabled(self.docx_btn, any(result.get("has_theft") for result in self.results))
            self.pages.setCurrentIndex(1)
        self._update_site_progress()
        QMessageBox.critical(self, "Extraction Power-Noc", message)

    def generate_pdf_reports(self):
        self._generate_reports(create_reports, "PDF")

    def generate_docx_reports(self):
        self._generate_reports(create_docx_reports, "Word")

    def _generate_reports(self, builder, kind):
        output_dir = QFileDialog.getExistingDirectory(
            self, f"Choisir le dossier des rapports {kind}"
        )
        if not output_dir:
            return

        previous_theme = self.chart._theme
        self.chart._theme = THEMES["light"]  # rapports lisibles à l'impression, quel que soit le thème
        try:
            paths = builder(
                self.results,
                output_dir,
                self.start_date.date().toString("yyyy-MM-dd"),
                self.end_date.date().toString("yyyy-MM-dd"),
                self._save_event_chart,
            )
        except PermissionError as exc:
            QMessageBox.critical(self, f"Rapports {kind}",
                                 f"Impossible d'écrire un rapport (est-il ouvert dans Word ou un lecteur PDF ?)\n{exc}")
            return
        except Exception as exc:
            QMessageBox.critical(self, f"Rapports {kind}", f"Impossible de générer les rapports : {exc}")
            return
        finally:
            # le graphique a servi de "brouillon" : on réaffiche la sélection courante
            self.chart._theme = previous_theme
            self.on_row_selected()
        self.status_label.setText(f"{len(paths)} rapport(s) {kind} généré(s)")
        if paths:
            QMessageBox.information(
                self,
                f"Rapports {kind}",
                f"{len(paths)} rapport(s) créé(s) dans :\n{output_dir}",
            )

    def _save_event_chart(self, result, output_path, event):
        self.chart.plot(result)
        left = max(0, event["row_start"] - 10)
        right = min(len(result["fuel"]) - 1, event["row_end"] + 10)
        self.chart.ax_fuel.set_xlim(left, max(left + 1, right))
        self.chart._rescale_y()
        theme = self.chart._theme
        annotate_drop(self.chart.ax_fuel, result["fuel"], event["row_start"], event["row_end"],
                      text_color=theme["text"], guide_color=theme["grid"], label_bg=theme["figure_bg"])
        self.chart.fig.savefig(output_path, dpi=160, bbox_inches="tight")

    def export_csv(self):
        path, _ = QFileDialog.getSaveFileName(self, "Exporter les résultats", "resultats_scan.csv", "CSV (*.csv)")
        if not path:
            return
        try:
            write_events_csv(path, self.results, mapping_key="site")
        except OSError as exc:
            QMessageBox.warning(self, "Export CSV", f"Impossible d'écrire le fichier (est-il ouvert dans Excel ?)\n{exc}")
            return
        self.status_label.setText(f"Données exportées : {os.path.basename(path)}")